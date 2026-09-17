import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import List
from app.exports.csv_export import flatten_org, FIELDNAMES


MINT = "FF14B8A6"
DARK = "FF334155"
LIGHT_GRAY = "FFF8FAFC"
WHITE = "FFFFFFFF"


def generate_excel(organizations: list, task_info: dict) -> bytes:
    wb = openpyxl.Workbook()

    # ─── Leads Sheet ─────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Leads"

    # Header row styling
    header_fill = PatternFill(start_color=DARK, end_color=DARK, fill_type="solid")
    header_font = Font(color=WHITE, bold=True, name="Calibri", size=11)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin = Side(border_style="thin", color="FFD1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Write headers
    for col_idx, field in enumerate(FIELDNAMES, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border

    ws.row_dimensions[1].height = 30

    # Data rows
    alt_fill = PatternFill(start_color="FFF1F5F9", end_color="FFF1F5F9", fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center", wrap_text=False)

    for row_idx, org in enumerate(organizations, start=2):
        flat = flatten_org(org)
        row_fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, field in enumerate(FIELDNAMES, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=flat.get(field, ""))
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = border
            if row_fill:
                cell.fill = row_fill

    # Color confidence column
    conf_col = FIELDNAMES.index("Confidence Level") + 1
    for row_idx in range(2, len(organizations) + 2):
        cell = ws.cell(row=row_idx, column=conf_col)
        if cell.value == "HIGH":
            cell.font = Font(name="Calibri", size=10, color="FF16A34A", bold=True)
        elif cell.value == "MEDIUM":
            cell.font = Font(name="Calibri", size=10, color="FFD97706", bold=True)
        else:
            cell.font = Font(name="Calibri", size=10, color="FFDC2626", bold=True)

    # Auto-fit columns
    for col_idx, field in enumerate(FIELDNAMES, start=1):
        max_len = len(field)
        for row_idx in range(2, min(len(organizations) + 2, 100)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, min(len(str(val)), 50))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    # Freeze header row
    ws.freeze_panes = "A2"

    # ─── Task Info Sheet ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Task Info")
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 40

    info_header_fill = PatternFill(start_color=MINT, end_color=MINT, fill_type="solid")
    title_cell = ws2.cell(row=1, column=1, value="LeadScrape Export")
    title_cell.font = Font(bold=True, size=14, color=WHITE)
    title_cell.fill = info_header_fill
    ws2.merge_cells("A1:B1")
    ws2.row_dimensions[1].height = 25

    info_rows = [
        ("Task ID", task_info.get("task_id", "")),
        ("Location", task_info.get("location", "")),
        ("Keyword", task_info.get("keyword", "")),
        ("Total Records", len(organizations)),
        ("Export Source", "LeadScrape v1.0"),
    ]
    for row_idx, (key, val) in enumerate(info_rows, start=2):
        ws2.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws2.cell(row=row_idx, column=2, value=str(val))

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
